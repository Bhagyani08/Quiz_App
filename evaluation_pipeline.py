#!/usr/bin/env python3
"""
End-to-End Quiz Evaluation Pipeline
Fetches bin IDs from admin bin → Downloads responses → Evaluates with LLM → Updates Excel
"""

import os
import sys
import json
import requests
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Add quiz_evaluation_system to path
sys.path.insert(0, 'quiz_evaluation_system')
import config
import llm_evaluator

# =============================================================
# CONFIGURATION
# =============================================================

JSONBIN_API_BASE = "https://api.jsonbin.io/v3"
ADMIN_BIN_ID = "6979e70643b1c97be951794c"

HEADERS = {
    "X-Master-Key": "$2a$10$0nEWKk89vS6CBYlIvV.zpuxU7Ja/DQ64Qk13e7mV60jM7ewVcYuGa",
    "Content-Type": "application/json"
}

# Output directory - now at project root level
OUTPUT_DIR = "evaluation_results"
EXCEL_FILE = f"{OUTPUT_DIR}/evaluation_report.xlsx"
JSON_FILE = f"{OUTPUT_DIR}/evaluation_results.json"

# =============================================================
# STEP 1: FETCH BIN IDS FROM ADMIN BIN
# =============================================================

def fetch_bin_ids_from_admin():
    """Fetch all bin IDs from admin bin"""
    print("📋 STEP 1: Fetching bin IDs from admin bin...")
    
    url = f"{JSONBIN_API_BASE}/b/{ADMIN_BIN_ID}/latest"
    
    try:
        response = requests.get(url, headers=HEADERS)
        if response.status_code == 200:
            admin_data = response.json()["record"]
            submissions = admin_data.get("submissions", [])
            bin_ids = [sub.get("bin_id") for sub in submissions if sub.get("bin_id")]
            
            print(f"   ✅ Found {len(bin_ids)} bin IDs in admin bin")
            return bin_ids
        else:
            print(f"   ❌ Failed to fetch admin bin: {response.status_code}")
            return []
    except Exception as e:
        print(f"   ❌ Error: {e}")
        return []

def get_processed_bin_ids():
    """Get list of bin IDs already processed in Excel"""
    if not os.path.exists(EXCEL_FILE):
        return set()
    
    try:
        df = pd.read_excel(EXCEL_FILE)
        return set(df['bin_id'].tolist())
    except Exception as e:
        print(f"   ⚠️  Could not read Excel file: {e}")
        return set()

# =============================================================
# STEP 2: DOWNLOAD BIN CONTENT
# =============================================================

def download_bin_content(bin_id):
    """Download content from a specific bin"""
    url = f"{JSONBIN_API_BASE}/b/{bin_id}/latest"
    
    try:
        response = requests.get(url, headers=HEADERS)
        if response.status_code == 200:
            return response.json()
        else:
            print(f"   ❌ Failed to download bin {bin_id}: {response.status_code}")
            return None
    except Exception as e:
        print(f"   ❌ Error downloading bin {bin_id}: {e}")
        return None

# =============================================================
# STEP 3: EVALUATE WITH LLM
# =============================================================

def extract_malpractice_data(response_data):
    """Extract malpractice information from response events"""
    events = response_data.get('record', {}).get('events', [])
    
    tab_switch_count = 0
    malpractice_detected = False
    exited_quiz = False
    
    for event in events:
        if event.get('event') == 'tab_switch':
            tab_switch_count = max(tab_switch_count, event.get('count', 0))
            
            warning = event.get('warning', '')
            if 'malpractice_detected' in warning:
                malpractice_detected = True
            if 'exited quiz' in warning.lower():
                exited_quiz = True
    
    if tab_switch_count >= 4:
        malpractice_detected = True
    
    return {
        'tab_switch_count': tab_switch_count,
        'malpractice_detected': malpractice_detected,
        'exited_quiz': exited_quiz
    }

def evaluate_bin(bin_id):
    """Download and evaluate a single bin"""
    print(f"\n   📥 Processing bin: {bin_id}")
    
    # Download content
    response_data = download_bin_content(bin_id)
    
    if not response_data:
        return None
    
    record = response_data.get('record', {})
    metadata = response_data.get('metadata', {})
    
    # Extract info
    name = record.get('name', 'Unknown')
    email = record.get('email', 'Unknown')
    answers = record.get('answers', [])
    
    print(f"      Student: {name}")
    print(f"      Questions: {len(answers)}")
    
    # Skip if no answers
    if len(answers) == 0:
        print(f"      ⚠️  Skipping: No answers")
        return None
    
    # Extract malpractice data
    malpractice_data = extract_malpractice_data(response_data)
    
    # Evaluate with LLM
    try:
        scores = llm_evaluator.evaluate_answers_with_llm(answers)
        score_summary = llm_evaluator.calculate_total_score(scores)
        
        print(f"      ✅ Score: {score_summary['total_score']}/{score_summary['max_score']} ({score_summary['percentage']}%)")
        
        evaluation = {
            'bin_id': bin_id,
            'name': name,
            'email': email,
            'scores': scores,
            **score_summary,
            **malpractice_data,
            'evaluation_timestamp': datetime.now().isoformat(),
            'created_at': metadata.get('createdAt', '')
        }
        
        return evaluation
        
    except Exception as e:
        print(f"      ❌ Evaluation failed: {e}")
        return {
            'bin_id': bin_id,
            'name': name,
            'email': email,
            'scores': [0.0] * len(answers),
            'total_score': 0.0,
            'max_score': float(len(answers)),
            'percentage': 0.0,
            **malpractice_data,
            'evaluation_timestamp': datetime.now().isoformat(),
            'created_at': metadata.get('createdAt', ''),
            'error': str(e)
        }

# =============================================================
# STEP 4: APPEND TO EXCEL
# =============================================================

def append_to_excel(new_evaluations):
    """Append new evaluations to existing Excel file"""
    print(f"\n📊 STEP 4: Appending {len(new_evaluations)} new evaluations to Excel...")
    
    if not new_evaluations:
        print("   ℹ️  No new evaluations to append")
        return
    
    # Convert to DataFrame
    rows = []
    for item in new_evaluations:
        row = {
            'bin_id': item.get('bin_id', ''),
            'name': item.get('name', ''),
            'email': item.get('email', ''),
            'total_score': item.get('total_score', 0),
            'max_score': item.get('max_score', 0),
            'percentage': item.get('percentage', 0),
            'tab_switch_count': item.get('tab_switch_count', 0),
            'malpractice_detected': item.get('malpractice_detected', False),
            'exited_quiz': item.get('exited_quiz', False),
            'evaluation_timestamp': item.get('evaluation_timestamp', ''),
            'created_at': item.get('created_at', ''),
        }
        
        if 'error' in item:
            row['error'] = item['error']
        
        rows.append(row)
    
    new_df = pd.DataFrame(rows)
    
    # Load existing Excel
    if os.path.exists(EXCEL_FILE):
        existing_df = pd.read_excel(EXCEL_FILE)
        
        # Remove duplicates based on bin_id
        existing_bin_ids = set(existing_df['bin_id'].tolist())
        new_df = new_df[~new_df['bin_id'].isin(existing_bin_ids)]
        
        if len(new_df) == 0:
            print("   ℹ️  All evaluations already exist in Excel")
            return
        
        # Combine
        combined_df = pd.concat([existing_df, new_df], ignore_index=True)
    else:
        combined_df = new_df
    
    # Sort by percentage
    combined_df = combined_df.sort_values('percentage', ascending=False)
    
    # Save to Excel
    combined_df.to_excel(EXCEL_FILE, index=False, sheet_name='Evaluation Results')
    
    # Apply formatting
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    
    # Red fill for malpractice
    red_fill = PatternFill(start_color='FFCCCC', end_color='FFCCCC', fill_type='solid')
    
    # Find malpractice column
    malpractice_col_idx = None
    for idx, cell in enumerate(ws[1], 1):
        if cell.value == 'malpractice_detected':
            malpractice_col_idx = idx
            break
    
    if malpractice_col_idx:
        for row_idx in range(2, ws.max_row + 1):
            malpractice_value = ws.cell(row=row_idx, column=malpractice_col_idx).value
            
            if malpractice_value == True or str(malpractice_value).lower() == 'true':
                for cell in ws[row_idx]:
                    cell.fill = red_fill
    
    # Auto-adjust columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(EXCEL_FILE)
    
    print(f"   ✅ Appended {len(new_df)} new evaluations")
    print(f"   📁 Excel updated: {EXCEL_FILE}")

def run_sample_json():
    ev = evaluate_bin("697b3d9a43b1c97be9548b47")

    print(ev)

# =============================================================
# MAIN PIPELINE
# =============================================================

def run_pipeline():
    """Run the complete evaluation pipeline"""
    print("=" * 80)
    print("🚀 Quiz Evaluation Pipeline - End to End Automation")
    print("=" * 80)
    print()
    
    # Step 1: Fetch bin IDs
    all_bin_ids = fetch_bin_ids_from_admin()
    
    if not all_bin_ids:
        print("\n❌ No bin IDs found. Exiting.")
        return
    
    # Filter out already processed bins
    print("\n🔍 STEP 1.5: Checking for already processed bins...")
    processed_bin_ids = get_processed_bin_ids()
    
    if processed_bin_ids:
        print(f"   ℹ️  Found {len(processed_bin_ids)} already processed bins")
        bin_ids = [bid for bid in all_bin_ids if bid not in processed_bin_ids]
        print(f"   ✅ {len(bin_ids)} new bins to process")
    else:
        bin_ids = all_bin_ids
        print(f"   ✅ All {len(bin_ids)} bins are new")
    
    if not bin_ids:
        print("\n✅ All bins already processed. Nothing to do!")
        return
    
    # Step 2 & 3: Download and evaluate each bin
    print(f"\n📊 STEP 2-3: Downloading and evaluating {len(bin_ids)} bins...")
    
    evaluations = []
    successful = 0
    failed = 0
    
    for idx, bin_id in enumerate(bin_ids, 1):
        print(f"\n[{idx}/{len(bin_ids)}]")
        
        evaluation = evaluate_bin(bin_id)
        
        if evaluation:
            evaluations.append(evaluation)
            if 'error' not in evaluation:
                successful += 1
            else:
                failed += 1
        else:
            failed += 1
    
    print(f"\n{'=' * 80}")
    print(f"Evaluation Summary:")
    print(f"   Total: {len(bin_ids)}")
    print(f"   Successful: {successful}")
    print(f"   Failed: {failed}")
    print(f"{'=' * 80}")
    
    # Create output directory
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    # Step 4: Append to Excel (only output format)
    append_to_excel(evaluations)
    
    print(f"\n✨ Pipeline completed successfully!")

if __name__ == "__main__":
    # run_sample_json()
    run_pipeline()
